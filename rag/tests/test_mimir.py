import os
import sys
import tempfile
import unittest
import shutil

# Ensure rag module is on sys.path
RAG_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
if RAG_DIR not in sys.path:
    sys.path.insert(0, RAG_DIR)

import db
from extractor import extract_text_from_file, extract_plain_text, extract_docx_text, extract_pdf_text
from indexer import index_files, index_folder, search_documents, get_embedding_model, get_vector_store
from llm import ask_question, query_ollama_direct


class TestExtractor(unittest.TestCase):
    """Tests for document text extraction across different file types."""

    def setUp(self):
        self.test_dir = tempfile.mkdtemp(prefix="mimir_test_extractor_")

    def tearDown(self):
        shutil.rmtree(self.test_dir, ignore_errors=True)

    def test_extract_plain_text_utf8(self):
        fpath = os.path.join(self.test_dir, "test.txt")
        with open(fpath, "w", encoding="utf-8") as f:
            f.write("Mimir AI local document indexing system.\nFast and private.")
        res = extract_text_from_file(fpath)
        self.assertEqual(res["type"], "txt")
        self.assertIn("Mimir AI local", res["text"])
        self.assertEqual(res["filename"], "test.txt")

    def test_extract_plain_text_utf16le(self):
        fpath = os.path.join(self.test_dir, "test_utf16.txt")
        with open(fpath, "w", encoding="utf-16-le") as f:
            f.write("Encoded in UTF-16 LE with BOM or without.")
        res = extract_text_from_file(fpath)
        self.assertIn("Encoded in UTF-16 LE", res["text"])

    def test_extract_markdown(self):
        fpath = os.path.join(self.test_dir, "guide.md")
        with open(fpath, "w", encoding="utf-8") as f:
            f.write("# Heading\n\n- Feature 1\n- Feature 2")
        res = extract_text_from_file(fpath)
        self.assertEqual(res["type"], "md")
        self.assertIn("Feature 1", res["text"])

    def test_extract_docx(self):
        import docx
        fpath = os.path.join(self.test_dir, "test.docx")
        doc = docx.Document()
        doc.add_paragraph("Paragraph 1 inside Word document.")
        doc.add_paragraph("Paragraph 2 with confidential metrics.")
        doc.save(fpath)

        res = extract_text_from_file(fpath)
        self.assertEqual(res["type"], "docx")
        self.assertIn("Paragraph 1 inside Word document.", res["text"])
        self.assertIn("confidential metrics", res["text"])

    def test_extract_pdf(self):
        import pymupdf
        fpath = os.path.join(self.test_dir, "test.pdf")
        doc = pymupdf.open()
        page = doc.new_page()
        page.insert_text(pymupdf.Point(50, 72), "Confidential quarterly revenue report: 4.2 million USD.")
        doc.save(fpath)
        doc.close()

        res = extract_text_from_file(fpath)
        self.assertEqual(res["type"], "pdf")
        self.assertIn("Confidential quarterly revenue report", res["text"])

    def test_nonexistent_file(self):
        res = extract_text_from_file(os.path.join(self.test_dir, "ghost_file.txt"))
        self.assertEqual(res["text"], "")
        self.assertIn("error", res)

    def test_empty_file(self):
        fpath = os.path.join(self.test_dir, "empty.txt")
        with open(fpath, "w", encoding="utf-8") as f:
            f.write("")
        res = extract_text_from_file(fpath)
        self.assertEqual(res["text"], "")


class TestDatabase(unittest.TestCase):
    """Tests for SQLite metadata storage."""

    def setUp(self):
        self.test_dir = tempfile.mkdtemp(prefix="mimir_test_db_")
        self.orig_db_path = db.DB_PATH
        db.DB_PATH = os.path.join(self.test_dir, "test_mimir.db")
        db.init_db()

    def tearDown(self):
        db.DB_PATH = self.orig_db_path
        shutil.rmtree(self.test_dir, ignore_errors=True)

    def test_init_and_stats_empty(self):
        stats = db.get_document_stats()
        self.assertEqual(stats["indexed_files"], 0)
        self.assertEqual(stats["total_size_bytes"], 0)
        self.assertEqual(stats["total_chunks"], 0)

    def test_add_and_retrieve_documents(self):
        db.add_or_update_document("resume.pdf", "/docs/resume.pdf", "pdf", 10240, 4)
        db.add_or_update_document("notes.txt", "/docs/notes.txt", "txt", 2048, 1)

        docs = db.get_all_documents()
        self.assertEqual(len(docs), 2)
        filenames = {d["filename"] for d in docs}
        self.assertIn("resume.pdf", filenames)
        self.assertIn("notes.txt", filenames)

        stats = db.get_document_stats()
        self.assertEqual(stats["indexed_files"], 2)
        self.assertEqual(stats["total_size_bytes"], 12288)
        self.assertEqual(stats["total_chunks"], 5)

    def test_update_existing_document(self):
        db.add_or_update_document("report.docx", "/docs/report.docx", "docx", 5000, 2)
        # Update same filepath with new size and chunk count
        db.add_or_update_document("report.docx", "/docs/report.docx", "docx", 8000, 3)

        docs = db.get_all_documents()
        self.assertEqual(len(docs), 1)
        self.assertEqual(docs[0]["file_size"], 8000)
        self.assertEqual(docs[0]["chunk_count"], 3)


class TestIndexerAndSearch(unittest.TestCase):
    """Tests for FAISS vector indexing and semantic retrieval."""

    def setUp(self):
        self.test_dir = tempfile.mkdtemp(prefix="mimir_test_index_")
        self.orig_index_dir = os.environ.get("INDEX_DIR")
        os.environ["INDEX_DIR"] = os.path.join(self.test_dir, "test_faiss")
        import indexer
        indexer.INDEX_DIR = os.environ["INDEX_DIR"]
        indexer._vector_store = None

    def tearDown(self):
        if self.orig_index_dir:
            os.environ["INDEX_DIR"] = self.orig_index_dir
        else:
            os.environ.pop("INDEX_DIR", None)
        import indexer
        indexer.INDEX_DIR = os.getenv("INDEX_DIR", os.path.join(RAG_DIR, "faiss_index"))
        indexer._vector_store = None
        shutil.rmtree(self.test_dir, ignore_errors=True)

    def test_embedding_model_shape(self):
        emb = get_embedding_model()
        vec = emb.embed_query("Local AI Assistant")
        self.assertEqual(len(vec), 384)

    def test_index_and_search_files(self):
        # Create 2 documents with distinct semantic topics
        doc1 = os.path.join(self.test_dir, "gardening.txt")
        with open(doc1, "w", encoding="utf-8") as f:
            f.write("Tomatoes require full sunlight and well-draining organic soil for optimal yield. Water deeply twice weekly.")

        doc2 = os.path.join(self.test_dir, "python_tips.txt")
        with open(doc2, "w", encoding="utf-8") as f:
            f.write("List comprehensions and generators provide concise, memory-efficient iterations in Python programming.")

        res = index_files([doc1, doc2])
        self.assertEqual(res["indexed_count"], 2)
        self.assertEqual(res["total_chunks"], 2)
        self.assertEqual(len(res["errors"]), 0)

        # Query semantic concepts
        results = search_documents("vegetable cultivation and watering vegetables", k=2)
        self.assertGreaterEqual(len(results), 1)
        top_result = results[0]
        self.assertEqual(top_result["filename"], "gardening.txt")
        self.assertIn("Tomatoes", top_result["snippet"])
        self.assertGreater(top_result["confidence"], 40)


class TestLLM(unittest.TestCase):
    """Tests for on-device Ollama model communication and grounded Q&A."""

    def setUp(self):
        self.test_dir = tempfile.mkdtemp(prefix="mimir_test_llm_")
        self.orig_index_dir = os.environ.get("INDEX_DIR")
        os.environ["INDEX_DIR"] = os.path.join(self.test_dir, "test_faiss")
        import indexer
        indexer.INDEX_DIR = os.environ["INDEX_DIR"]
        indexer._vector_store = None

        # Create and index a test document for grounded Q&A
        doc_path = os.path.join(self.test_dir, "internship_offer.txt")
        with open(doc_path, "w", encoding="utf-8") as f:
            f.write("Internship Offer Letter: The notice period for termination is two weeks. Monthly stipend is 4500 dollars.")
        index_files([doc_path])

    def tearDown(self):
        if self.orig_index_dir:
            os.environ["INDEX_DIR"] = self.orig_index_dir
        else:
            os.environ.pop("INDEX_DIR", None)
        import indexer
        indexer.INDEX_DIR = os.getenv("INDEX_DIR", os.path.join(RAG_DIR, "faiss_index"))
        indexer._vector_store = None
        shutil.rmtree(self.test_dir, ignore_errors=True)

    def test_ollama_direct_connection(self):
        messages = [{"role": "user", "content": "Respond with the single word: OK"}]
        resp = query_ollama_direct("qwen2.5:3b", messages, timeout=15)
        self.assertIsNotNone(resp)
        self.assertIn("OK", resp.upper())

    def test_ask_question_grounded(self):
        res = ask_question("What is the notice period mentioned in the offer letter?")
        self.assertIn("answer", res)
        self.assertIn("sources", res)
        self.assertGreaterEqual(len(res["sources"]), 1)
        self.assertTrue(
            "two weeks" in res["answer"].lower() or "2 weeks" in res["answer"].lower(),
            f"Unexpected answer: {res['answer']}"
        )


class TestFastAPIEndpoints(unittest.TestCase):
    """Tests for all REST API endpoints."""

    @classmethod
    def setUpClass(cls):
        from fastapi.testclient import TestClient
        from app import app
        cls.client = TestClient(app)

    def test_health_check(self):
        res = self.client.get("/health")
        self.assertEqual(res.status_code, 200)
        data = res.json()
        self.assertEqual(data["status"], "healthy")
        self.assertTrue(data["on_device"])
        self.assertIn("stats", data)

    def test_get_documents(self):
        res = self.client.get("/documents")
        self.assertEqual(res.status_code, 200)
        data = res.json()
        self.assertEqual(data["status"], "success")
        self.assertIn("documents", data)
        self.assertIn("stats", data)

    def test_search_endpoint(self):
        res = self.client.post("/search", json={"query": "internship stipend", "k": 3})
        self.assertEqual(res.status_code, 200)
        data = res.json()
        self.assertEqual(data["status"], "success")
        self.assertIn("results", data)

    def test_ask_endpoint(self):
        res = self.client.post("/ask", json={"query": "What is the monthly stipend?"})
        self.assertEqual(res.status_code, 200)
        data = res.json()
        self.assertEqual(data["status"], "success")
        self.assertIn("answer", data)
        self.assertIn("sources", data)

    def test_upload_endpoint(self):
        # Create a mock file in memory and upload
        file_content = b"Local file upload test content for on-device vector indexing."
        files = {
            "files": ("uploaded_test_file.txt", file_content, "text/plain")
        }
        res = self.client.post("/upload", files=files)
        self.assertEqual(res.status_code, 200)
        data = res.json()
        self.assertEqual(data["status"], "success")
        self.assertGreaterEqual(data["data"]["indexed_count"], 1)


if __name__ == "__main__":
    unittest.main()
